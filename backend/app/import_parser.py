"""Bounded, deterministic import validation; no database or network access."""

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook

FINANCIAL_COLUMNS = (
    "period",
    "department",
    "account_code",
    "account_name",
    "account_type",
    "amount",
)
CASH_COLUMNS = ("expected_date", "description", "category", "direction", "amount", "status")
ACCOUNT_TYPES = {"revenue", "cogs", "operating_expense", "other_income", "other_expense"}
MAX_ROWS = 50000
MAX_COLUMNS = 64


@dataclass
class ParsedImport:
    rows: list[dict]
    warnings: list[dict]


class ImportValidationError(ValueError):
    def __init__(self, issues: list[dict]):
        self.issues = issues[:100]
        super().__init__("The file contains invalid rows. Correct the errors and upload it again.")


def problem(row: int, field: str, message: str) -> dict:
    return {"row": row, "field": field, "message": message}


def source_rows(data: bytes, filename: str):
    if filename.endswith(".csv"):
        try:
            stream = io.StringIO(data.decode("utf-8-sig"), newline="")
            yield from csv.reader(stream, strict=True)
        except (UnicodeDecodeError, csv.Error) as error:
            raise ImportValidationError(
                [problem(1, "file", "Use a valid UTF-8 CSV file.")]
            ) from error
        return
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 1000
                or sum(entry.file_size for entry in entries) > 100 * 1024 * 1024
                or any("vbaproject" in entry.filename.lower() for entry in entries)
            ):
                raise ValueError("Unsafe workbook")
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            sheet = workbook.worksheets[0]
            sheet.reset_dimensions()
            for number, cells in enumerate(sheet.iter_rows(), 1):
                if len(cells) > MAX_COLUMNS:
                    raise ValueError("Too many columns")
                if any(cell.data_type == "f" for cell in cells):
                    raise ImportValidationError(
                        [problem(number, "file", "Replace formulas with literal values.")]
                    )
                yield [cell.value for cell in cells]
        finally:
            workbook.close()
    except ImportValidationError:
        raise
    except (
        BadZipFile,
        ValueError,
        KeyError,
        IndexError,
        OSError,
        TypeError,
        ParseError,
        DefusedXmlException,
    ) as error:
        raise ImportValidationError(
            [problem(1, "file", "Use a valid XLSX workbook without macros or formulas.")]
        ) from error


def parse_date(value, monthly: bool) -> str:
    if isinstance(value, datetime):
        value = value.date()
    if not isinstance(value, date):
        text = str(value).strip()
        if not re.fullmatch(r"\d{4}-\d{2}(-\d{2})?" if monthly else r"\d{4}-\d{2}-\d{2}", text):
            raise ValueError("Invalid date format")
        if monthly and len(text) == 7:
            text += "-01"
        value = date.fromisoformat(text)
    return (value.replace(day=1) if monthly else value).isoformat()


def parse_import(data: bytes, filename: str, kind: str) -> ParsedImport:
    required = CASH_COLUMNS if kind == "cash" else FINANCIAL_COLUMNS
    source = iter(source_rows(data, filename))
    header = next(source, [])
    header = [str(value or "").strip() for value in header]
    if len(header) > MAX_COLUMNS or len(set(header)) != len(header):
        raise ImportValidationError([problem(1, "header", "Column headings must be unique.")])
    missing = set(required) - set(header)
    if missing:
        raise ImportValidationError(
            [problem(1, "header", "Missing columns: " + ", ".join(sorted(missing)))]
        )
    warnings = []
    if set(header) - set(required):
        warnings.append(problem(1, "header", "Extra columns were ignored."))
    rows, issues = [], []
    for number, values in enumerate(source, 2):
        if number > MAX_ROWS + 1:
            raise ImportValidationError([problem(number, "file", "Maximum 50,000 data rows.")])
        if not any(value is not None and str(value).strip() for value in values):
            continue
        if len(values) != len(header):
            issues.append(problem(number, "row", "The row does not match the column headings."))
            if len(issues) >= 100:
                break
            continue
        row = {}
        raw = dict(zip(header, values, strict=True))
        for field in required:
            value = raw[field]
            try:
                if value is None or not str(value).strip():
                    raise ValueError("A value is required.")
                if field in {"period", "expected_date"}:
                    row[field] = parse_date(value, field == "period")
                elif field == "amount":
                    if isinstance(value, bool):
                        raise ValueError("Use a finite number without currency symbols.")
                    amount = Decimal(str(value).strip())
                    if not amount.is_finite() or len(amount.as_tuple().digits) > 30:
                        raise ValueError("Use a finite number with at most 30 significant digits.")
                    if abs(amount) > Decimal("1e20") or amount.as_tuple().exponent < -10:
                        raise ValueError("Amount is outside the supported range or precision.")
                    if kind == "cash" and amount < 0:
                        raise ValueError("Cash amounts must be nonnegative; use direction.")
                    row[field] = str(amount)
                    if amount < 0 and len(warnings) < 100:
                        warnings.append(problem(number, field, "Signed financial amount retained."))
                else:
                    text = str(value).strip()
                    if len(text) > 500:
                        raise ValueError("Maximum 500 characters.")
                    allowed = {
                        "account_type": ACCOUNT_TYPES,
                        "direction": {"inflow", "outflow"},
                        "status": {"planned", "confirmed", "actual"},
                    }.get(field)
                    if allowed and text not in allowed:
                        raise ValueError("Use one of: " + ", ".join(sorted(allowed)))
                    row[field] = text
            except (ValueError, InvalidOperation) as error:
                message = (
                    str(error)
                    if field not in {"period", "expected_date", "amount"}
                    else {
                        "period": "Use a valid YYYY-MM or YYYY-MM-DD date.",
                        "expected_date": "Use a valid YYYY-MM-DD date.",
                        "amount": "Use a finite numeric amount; cash must be nonnegative.",
                    }[field]
                )
                issues.append(problem(number, field, message))
        row["source_row"] = number
        rows.append(row)
        if len(issues) >= 100:
            break
    if issues:
        raise ImportValidationError(issues)
    if not rows:
        raise ImportValidationError([problem(2, "file", "Add at least one data row.")])
    return ParsedImport(rows, warnings)
