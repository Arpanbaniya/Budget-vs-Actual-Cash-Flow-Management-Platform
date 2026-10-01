"""Private Excel reporting with owner-scoped metadata and expiring downloads."""

import asyncio
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Annotated
from urllib.parse import parse_qs, urlsplit
from uuid import UUID, uuid4

from fastapi import APIRouter, Depends, Query, Response
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from starlette.concurrency import run_in_threadpool

from app.cash import owned_record, record_scope
from app.companies import Company, Store
from app.data import exact_json, read_rows
from app.errors import ApiError
from app.forecast import forecast_data
from app.imports import private_response
from app.insights import AnalysisRequest
from app.variance import calculate_variance, financial_data

router = APIRouter(prefix="/api/v1", tags=["reports"], dependencies=[Depends(private_response)])
BUCKET = "fpna-reports"
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEETS = [
    "Executive Summary",
    "Variance Analysis",
    "Department Analysis",
    "Cash Forecast",
    "Scenario",
    "Source Budget",
    "Source Actual",
    "Source Cash",
]


def safe_text(value: str) -> str:
    text = ILLEGAL_CHARACTERS_RE.sub("", value)[:32766]
    return (
        "'" + text
        if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n"))
        else text
    )


def excel_value(value):
    if value is None:
        return None
    if isinstance(value, Decimal):
        # Excel has 15-digit numeric precision. Preserve larger decimals as text.
        return value if len(value.as_tuple().digits) <= 15 else str(value)
    if isinstance(value, str):
        return safe_text(value)
    return value


def build_workbook(
    company: Company,
    body: AnalysisRequest,
    lines: list[dict],
    cash_items: list[dict],
    forecast: dict,
    generated_at: datetime,
) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    stamp = generated_at.isoformat()

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        ws.append([safe_text(company.name)])
        ws.append([f"Period: {body.start} to {body.end}"])
        ws.append([f"Currency: {company.currency}", f"Generated: {stamp}"])
        ws.append(["Amounts exceeding Excel's 15-digit precision are stored as text."])
        ws.append(headers)
        for row in rows:
            ws.append([excel_value(value) for value in row])
        ws.freeze_panes = "A6"
        ws.auto_filter.ref = f"A5:{ws.cell(ws.max_row, len(headers)).coordinate}"
        for cell in ws[5]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="164D3B")
        ws.row_dimensions[5].height = 25
        for column in ws.columns:
            letter = column[0].column_letter
            ws.column_dimensions[letter].width = min(
                48, max(18, max(len(str(cell.value or "")) for cell in column[4:]) + 2)
            )
        for row in ws.iter_rows(min_row=6):
            for cell in row:
                if cell.data_type == "n":
                    cell.number_format = "#,##0.00;[Red](#,##0.00)"
        return ws

    variance = calculate_variance(lines)
    summary = variance["summary"]
    sheet(
        "Executive Summary",
        ["Metric", "Budget / value", "Actual", "Variance"],
        [
            *[
                [
                    metric.replace("_", " ").title(),
                    summary["budget_" + metric],
                    summary["actual_" + metric],
                    summary[metric + "_variance"],
                ]
                for metric in ("revenue", "expenses", "operating_profit")
            ],
            ["Opening cash", forecast["opening_cash"]],
            ["Opening snapshot date", forecast["balance_date"]],
            ["Forecast start", body.cash_start_date],
            ["Minimum 13-week closing cash", forecast["minimum_projected_cash"]],
            ["First threshold breach week", forecast["first_threshold_breach_week"]],
            ["Minimum cash threshold", forecast["minimum_cash_threshold"]],
            ["Budget source rows", variance["budget_row_count"]],
            ["Actual source rows", variance["actual_row_count"]],
            [
                "Snapshot method",
                "Opening balance used directly; earlier items are not rolled forward.",
            ],
        ],
    )
    columns = [
        "Group",
        "Account name",
        "Account type",
        "Budget",
        "Actual",
        "Variance",
        "Variance %",
        "Result",
    ]
    for name, group in [("Variance Analysis", "account"), ("Department Analysis", "department")]:
        rows = calculate_variance(lines, group)["rows"]
        sheet(
            name,
            columns,
            [
                [
                    row[key]
                    for key in (
                        "label",
                        "account_name",
                        "account_type",
                        "budget_amount",
                        "actual_amount",
                        "variance_amount",
                        "variance_percent",
                        "favorability",
                    )
                ]
                for row in rows
            ],
        )
    sheet(
        "Cash Forecast",
        ["Week", "Start", "End", "Opening", "Inflows", "Outflows", "Closing", "Below threshold"],
        [
            [
                row[key]
                for key in (
                    "week_number",
                    "week_start",
                    "week_end",
                    "opening_cash",
                    "inflows",
                    "outflows",
                    "closing_cash",
                    "threshold_breached",
                )
            ]
            for row in forecast["weekly"]
        ],
    )
    scenario = forecast["scenario"]
    sheet(
        "Scenario",
        ["Setting", "Value"],
        [
            ["Name", scenario["name"] if scenario else "Base forecast"],
            [
                "Inflow adjustment %",
                Decimal(str(scenario["inflow_adjustment_pct"])) if scenario else Decimal(0),
            ],
            [
                "Outflow adjustment %",
                Decimal(str(scenario["outflow_adjustment_pct"])) if scenario else Decimal(0),
            ],
            ["Collection delay days", scenario["collection_delay_days"] if scenario else 0],
            [
                "Rules",
                "Adjustments and collection delays apply only to planned items. Source rows are unchanged.",
            ],
        ],
    )
    for name, kind in [("Source Budget", "budget"), ("Source Actual", "actual")]:
        sheet(
            name,
            [
                "Period",
                "Department",
                "Account code",
                "Account name",
                "Account type",
                "Amount",
                "Import ID",
                "Source row",
            ],
            [
                [
                    row["period"],
                    row["department"],
                    row["account_code"],
                    row["account_name"],
                    row["account_type"],
                    Decimal(str(row["amount"])),
                    row.get("import_id"),
                    row.get("source_row"),
                ]
                for row in lines
                if row["kind"] == kind
            ],
        )
    sheet(
        "Source Cash",
        [
            "Expected date",
            "Description",
            "Category",
            "Direction",
            "Amount",
            "Status",
            "Import ID",
            "Source row",
        ],
        [
            [
                row["expected_date"],
                row["description"],
                row["category"],
                row["direction"],
                Decimal(str(row["amount"])),
                row["status"],
                row.get("import_id"),
                row.get("source_row"),
            ]
            for row in cash_items
        ],
    )
    output = BytesIO()
    try:
        wb.save(output)
    finally:
        wb.close()
    return output.getvalue()


async def owned_report(store, report_id):
    record = await owned_record(store, "reports", report_id)
    expected = f"{store.user.user_id}/{record['company_id']}/{record['id']}/report.xlsx"
    if record.get("storage_path") not in {None, expected}:
        raise ApiError(409, "REPORT_INVALID", "The report storage path is invalid.")
    return record


@router.post("/companies/{company_id}/reports/excel", status_code=201)
async def generate_report(company_id: UUID, body: AnalysisRequest, store: Store):
    company = await store.get(company_id)
    # Validate the forecast before creating metadata, including owned scenario and snapshot.
    lines, forecast = await asyncio.gather(
        financial_data(store, company_id, body.start, body.end),
        forecast_data(store, company, body.cash_start_date, 13, body.scenario_id),
    )
    delay = forecast["scenario"]["collection_delay_days"] if forecast["scenario"] else 0
    source_start = body.cash_start_date - timedelta(
        days=min(delay, (body.cash_start_date - date.min).days)
    )
    cash_items = await read_rows(
        store,
        "cash_items",
        company_id,
        **{
            "and": f"(expected_date.gte.{source_start},expected_date.lte.{body.cash_start_date + timedelta(days=90)})"
        },
    )
    if len(lines) + len(cash_items) > 100000:
        raise ApiError(
            422,
            "REPORT_LIMIT",
            "Narrow the report period to export fewer than 100,000 source rows.",
        )
    id = uuid4()
    path = f"{store.user.user_id}/{company_id}/{id}/report.xlsx"
    await store.request(
        "POST",
        "rest/v1/reports",
        json={
            "id": str(id),
            "user_id": store.user.user_id,
            "company_id": str(company_id),
            "status": "generating",
            "storage_path": path,
            "parameters": exact_json(body.model_dump(by_alias=True)),
        },
        headers={"Prefer": "return=representation"},
    )
    uploaded = False
    try:
        content = await run_in_threadpool(
            build_workbook, company, body, lines, cash_items, forecast, datetime.now(UTC)
        )
        if len(content) > 20 * 1024 * 1024:
            raise ApiError(422, "REPORT_LIMIT", "Narrow the period to create a smaller report.")
        uploaded = True
        await store.request(
            "POST",
            f"storage/v1/object/{BUCKET}/{path}",
            content=content,
            headers={"Content-Type": MIME, "x-upsert": "false"},
        )
        response = await store.request(
            "PATCH",
            "rest/v1/reports",
            params=record_scope(store, id),
            json={
                "status": "ready",
                "completed_at": datetime.now(UTC).isoformat(),
                "error_message": None,
            },
            headers={"Prefer": "return=representation"},
        )
        return exact_json(store.rows(response)[0])
    except Exception as error:
        if uploaded:
            try:
                await store.request(
                    "DELETE", f"storage/v1/object/{BUCKET}", json={"prefixes": [path]}
                )
            except ApiError:
                pass
        try:
            await store.request(
                "PATCH",
                "rest/v1/reports",
                params=record_scope(store, id),
                json={
                    "status": "failed",
                    "error_message": "Report generation failed. Delete this entry and generate a new report.",
                },
            )
        except ApiError:
            pass
        if isinstance(error, ApiError):
            raise
        raise ApiError(503, "REPORT_UNAVAILABLE", "Report generation failed. Try again.") from error


@router.get("/companies/{company_id}/reports")
async def list_reports(
    company_id: UUID,
    store: Store,
    page: Annotated[int, Query(ge=1)] = 1,
    page_size: Annotated[int, Query(ge=1, le=100)] = 20,
):
    await store.get(company_id)
    response = await store.request(
        "GET",
        "rest/v1/reports",
        params={
            "user_id": f"eq.{store.user.user_id}",
            "company_id": f"eq.{company_id}",
            "select": "*",
            "order": "created_at.desc,id.desc",
            "offset": str((page - 1) * page_size),
            "limit": str(page_size),
        },
    )
    return exact_json({"items": store.rows(response), "page": page, "page_size": page_size})


@router.get("/reports/{report_id}")
async def get_report(report_id: UUID, store: Store):
    record = await owned_report(store, report_id)
    if record["status"] != "ready":
        return exact_json(record)
    path = record["storage_path"]
    if not path:
        raise ApiError(409, "REPORT_INVALID", "The report is missing its file.")
    response = await store.request(
        "POST", f"storage/v1/object/sign/{BUCKET}/{path}", json={"expiresIn": 300}
    )
    relative = response.json().get("signedURL")
    if not isinstance(relative, str):
        raise ApiError(503, "REPORT_UNAVAILABLE", "A download link could not be created.")
    parsed = urlsplit(relative)
    if (
        parsed.scheme
        or parsed.netloc
        or parsed.fragment
        or parsed.path != f"/object/sign/{BUCKET}/{path}"
        or not parse_qs(parsed.query).get("token")
    ):
        raise ApiError(503, "REPORT_UNAVAILABLE", "A download link could not be created.")
    return exact_json(
        {
            **record,
            "download_url": f"{str(store.client.base_url).rstrip('/')}/storage/v1{relative}",
            "expires_in_seconds": 300,
        }
    )


@router.delete("/reports/{report_id}", status_code=204)
async def delete_report(report_id: UUID, store: Store):
    record = await owned_report(store, report_id)
    if record["status"] == "generating" and datetime.now(UTC) - datetime.fromisoformat(
        record["created_at"]
    ) < timedelta(minutes=10):
        raise ApiError(409, "REPORT_BUSY", "The report is still generating. Try again later.")
    if record.get("storage_path"):
        await store.request(
            "DELETE", f"storage/v1/object/{BUCKET}", json={"prefixes": [record["storage_path"]]}
        )
    await store.request("DELETE", "rest/v1/reports", params=record_scope(store, report_id))
    return Response(status_code=204, headers={"Cache-Control": "private, no-store"})
