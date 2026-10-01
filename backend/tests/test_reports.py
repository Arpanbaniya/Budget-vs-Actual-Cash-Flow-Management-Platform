from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from test_imports import COMPANY_A, auth
from test_variance import line

from app.companies import Company
from app.forecast import calculate_forecast
from app.insights import AnalysisRequest
from app.reports import SHEETS, build_workbook

PARAMS = {"from": "2026-10-01", "to": "2026-10-31", "cash_start_date": "2026-10-01"}


def test_workbook_sheets_cells_precision_and_sanitization(imports_api):
    company = Company.model_validate(
        {**imports_api["companies"][COMPANY_A], "name": '=HYPERLINK("bad")'}
    )
    lines = [
        line("budget", "100"),
        {**line("actual", "90"), "account_name": "+cmd", "department": "\t=evil"},
        {
            **line("actual", "100000000000000000001", "cogs"),
            "account_name": "@evil",
            "account_code": "0002",
        },
    ]
    forecast = calculate_forecast(
        {"amount": "200", "balance_date": "2026-10-01"}, [], date(2026, 10, 1), 13, Decimal(50)
    )
    content = build_workbook(
        company, AnalysisRequest.model_validate(PARAMS), lines, [], forecast, datetime.now(UTC)
    )
    wb = load_workbook(BytesIO(content), data_only=False)
    assert wb.sheetnames == SHEETS
    assert wb["Executive Summary"]["C6"].value == 90
    assert wb["Executive Summary"]["D6"].value == -10
    assert wb["Cash Forecast"]["G6"].value == 200
    assert wb["Source Actual"]["C7"].value == "0002"
    assert wb["Source Actual"]["F7"].value == "100000000000000000001"
    assert wb["Source Actual"]["D6"].value.startswith("'")
    for ws in wb:
        assert ws.freeze_panes == "A6"
        assert all(cell.data_type != "f" for row in ws for cell in row)
        assert ws["A1"].value.startswith("'")
    wb.close()


def create_report(env):
    base = f"/api/v1/companies/{COMPANY_A}"
    env["client"].post(
        base + "/cash-balances",
        headers=auth(),
        json={"balance_date": "2026-10-01", "amount": "200"},
    )
    return env["client"].post(base + "/reports/excel", headers=auth(), json=PARAMS)


def test_report_private_upload_signed_download_and_delete(imports_api):
    env = imports_api
    response = create_report(env)
    assert response.status_code == 201
    row = response.json()
    assert row["status"] == "ready"
    assert "download_url" not in row
    wb = load_workbook(BytesIO(env["objects"][row["storage_path"]]))
    assert wb.sheetnames == SHEETS
    wb.close()
    path = "/api/v1/reports/" + row["id"]
    download = env["client"].get(path, headers=auth())
    assert download.json()["expires_in_seconds"] == 300
    assert download.json()["download_url"].startswith(
        "https://example.supabase.co/storage/v1/object/sign/fpna-reports/"
    )
    for method in ("get", "delete"):
        assert getattr(env["client"], method)(path, headers=auth("user-b")).status_code == 404
    assert env["client"].delete(path, headers=auth()).status_code == 204
    assert row["storage_path"] not in env["objects"] and not env["reports"]


@pytest.mark.parametrize(
    "url",
    [
        "https://evil.example/?token=x",
        "/object/sign/fpna-reports/other/report.xlsx?token=x",
        "/object/sign/fpna-reports/other/report.xlsx",
    ],
)
def test_report_rejects_invalid_signed_urls(imports_api, url):
    row = create_report(imports_api).json()
    imports_api["controls"]["report_signed_url"] = url
    response = imports_api["client"].get("/api/v1/reports/" + row["id"], headers=auth())
    assert response.status_code == 503
    assert url not in response.text


def test_report_upload_failure_keeps_failed_metadata(imports_api):
    env = imports_api
    # Inject failure after metadata creation, when upload starts.
    import app.reports as module

    original = module.build_workbook

    def build(*args):
        path = env["reports"][0]["storage_path"]
        env["failures"][f"POST /storage/v1/object/fpna-reports/{path}"] = 503
        return original(*args)

    from unittest.mock import patch

    with patch.object(module, "build_workbook", build):
        assert create_report(env).status_code == 503
    assert env["reports"][0]["status"] == "failed"
    assert not env["objects"]
